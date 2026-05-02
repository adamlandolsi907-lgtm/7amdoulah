"""
Excel operational report extractor for ADWYA tri-generation BILAN TOTAL sheets.

Sheet structure (monthly time-series log, ~10-min intervals):
  - Col A: Category label (merged cells, changes at section boundaries)
  - Col B: Measure description
  - Cols D+: Chronological readings (cumulative meter values)

Strategy: compute DELTA = last_value - first_value for cumulative meters.
Emits one EnergyRecord per energy stream per file.
"""

import re
import uuid
from datetime import datetime
from pathlib import Path
from typing import Optional, Union

import openpyxl

from app.models.schemas import DocumentType, EnergyRecord, EnergyType

AnyRecord = EnergyRecord

# ── Row targets ───────────────────────────────────────────────────────────────
# Each entry: (category_re, label_re, field_name, unit_raw, energy_type)
# category_re matched against col-A label (accumulated downward)
# label_re matched against col-B label
_TARGETS: list[tuple[str, str, str, str, EnergyType]] = [
    (r"consommation gaz",             r"gaz naturel moteur.*nm3",          "gas_motor_nm3",               "Nm3",  EnergyType.natural_gas),
    (r"consommation auxiliar",        r"energie.*electr.*kwh",              "aux_electricity_kwh",         "kWh",  EnergyType.electricity),
    (r"energie moteur",               r"alternateur.*kwh",                  "electricity_generated_kwh",   "kWh",  EnergyType.electricity),
    (r"energymeter eau glac",         r"energie en kwh",                    "cooling_kwh",                 "kWh",  EnergyType.chilled_water),
    (r"energymeter eau chaude r[eé]cup", r"energie en kwh",                 "hot_water_recovered_kwh",     "kWh",  EnergyType.hot_water),
    (r"energymeter eau chaude.*alpha sanitaire", r"energie en kwh",         "hot_water_alpha_sanitary_kwh","kWh",  EnergyType.hot_water),
    (r"energymeter eau chaude alpha$",r"energie en kwh",                    "hot_water_alpha_kwh",         "kWh",  EnergyType.hot_water),
    (r"energymeter eau chaude gamma", r"energie en kwh",                    "hot_water_gamma_kwh",         "kWh",  EnergyType.hot_water),
    (r"achat et vente",               r"energie positive steg",             "grid_purchase_kwh",           "kWh",  EnergyType.electricity),
    (r"achat et vente",               r"energie negative steg",             "grid_injection_kwh",          "kWh",  EnergyType.electricity),
    (r"achat et vente",               r"energie positive production",       "production_kwh",              "kWh",  EnergyType.electricity),
]

# Zone hint from field name
_ZONE_MAP = {
    "hot_water_alpha_sanitary_kwh":  "alpha",
    "hot_water_alpha_kwh":           "alpha",
    "hot_water_gamma_kwh":           "gamma",
    "grid_injection_kwh":            "grid_export",
    "grid_purchase_kwh":             "grid_import",
    "electricity_generated_kwh":     "self_generation",
    "aux_electricity_kwh":           "self_generation",
}


# ── Date helpers ──────────────────────────────────────────────────────────────

def _parse_cell_date(val) -> Optional[str]:
    """Convert a date cell value (datetime object or date string) to YYYY-MM."""
    if isinstance(val, datetime):
        return val.strftime("%Y-%m")
    if isinstance(val, str):
        # ISO format: YYYY-MM-DD
        m = re.match(r"(20\d{2})[/\-](\d{2})[/\-]\d{2}", val)
        if m:
            return f"{m.group(1)}-{m.group(2)}"
        # DD/MM/YYYY or DD-MM-YYYY
        m = re.search(r"(\d{1,2})[/\-](\d{1,2})[/\-](20\d{2})", val)
        if m:
            return f"{m.group(3)}-{m.group(2).zfill(2)}"
    return None


_MONTH_FR = {
    "janvier": "01", "fevrier": "02", "février": "02", "mars": "03",
    "avril": "04", "mai": "05", "juin": "06", "juillet": "07",
    "aout": "08", "août": "08", "septembre": "09", "octobre": "10",
    "november": "11", "novembre": "11", "decembre": "12", "décembre": "12",
}
_MONTH_EN = {
    "january": "01", "february": "02", "march": "03", "april": "04",
    "may": "05", "june": "06", "july": "07", "august": "08",
    "september": "09", "october": "10", "november": "11", "december": "12",
}


def _date_from_filename(name: str) -> Optional[str]:
    """Attempt YYYY-MM extraction from filename like 'FIN AOUT 2025' or 'avril-report1'."""
    lower = name.lower()
    for word, mo in {**_MONTH_FR, **_MONTH_EN}.items():
        if word in lower:
            year_m = re.search(r"(20\d{2})", name)
            year = year_m.group(1) if year_m else None
            if not year:
                # Try the _DDMMYYYY suffix: e.g. _2292025 → year=2025
                suffix = re.search(r"_\d{1,2}\d{2}(20\d{2})", name)
                year = suffix.group(1) if suffix else None
            if year:
                return f"{year}-{mo}"
    return None


# ── Row scanner ───────────────────────────────────────────────────────────────

def _scan_rows(ws) -> dict[str, list]:
    """
    Walk rows; for each target, record the list of numeric data values
    from columns D onwards (index 3+).
    Returns {field_name: [v1, v2, ..., vN]} with only non-None numerics.
    Also returns special key "_date_row" with all date-column values.
    """
    current_cat = ""
    result: dict[str, list] = {t[2]: [] for t in _TARGETS}
    result["_date_row"] = []

    compiled = [
        (re.compile(cat_re, re.I), re.compile(lbl_re, re.I), fname)
        for cat_re, lbl_re, fname, *_ in _TARGETS
    ]

    for row in ws.iter_rows(values_only=True):
        a_val = str(row[0] or "").strip()
        b_val = str(row[1] or "").strip()

        if a_val:
            current_cat = a_val

        # Check if this is the date header row
        if b_val.lower() == "date":
            for cell in row[3:]:
                if cell is not None:
                    result["_date_row"].append(cell)
            continue

        # Check each target
        for cat_re, lbl_re, fname in compiled:
            if not cat_re.search(current_cat):
                continue
            if not lbl_re.search(b_val):
                continue
            # Extract numeric values from data columns (D onwards, index 3+)
            for cell in row[3:]:
                v = _to_num(cell)
                if v is not None:
                    result[fname].append(v)
            break

    return result


def _to_num(val) -> Optional[float]:
    if val is None:
        return None
    if isinstance(val, (int, float)):
        return float(val)
    try:
        return float(str(val).replace(" ", "").replace(",", "."))
    except (ValueError, TypeError):
        return None


def _delta(values: list[float]) -> Optional[float]:
    """Return last - first for a cumulative meter reading list."""
    nums = [v for v in values if v is not None]
    if len(nums) < 2:
        return nums[0] if nums else None
    return round(nums[-1] - nums[0], 4)


# ── Public entry point ────────────────────────────────────────────────────────

def extract_excel(excel_path: str) -> list[EnergyRecord]:
    """
    Extract monthly energy stream records from a BILAN TOTAL Excel file.
    Returns one EnergyRecord per energy stream with non-zero delta.
    """
    source_file = Path(excel_path).name
    wb = openpyxl.load_workbook(excel_path, data_only=True)

    # Try each sheet; prefer BILAN TOTAL
    ws = None
    for name in wb.sheetnames:
        if "bilan" in name.lower() or "total" in name.lower():
            ws = wb[name]
            break
    if ws is None:
        ws = wb.active

    scanned = _scan_rows(ws)
    wb.close()

    # Determine reporting date
    date_vals = scanned.pop("_date_row", [])
    date_str = None
    for dv in reversed(date_vals):          # last timestamp = end of period
        date_str = _parse_cell_date(dv)
        if date_str:
            break
    if not date_str:
        date_str = _date_from_filename(source_file) or "unknown"

    # Build EnergyRecord per target
    records: list[EnergyRecord] = []
    target_meta = {t[2]: (t[3], t[4]) for t in _TARGETS}

    for field_name, values in scanned.items():
        delta = _delta(values)
        if delta is None or delta == 0.0:
            continue

        unit_raw, energy_type = target_meta[field_name]
        zone = _ZONE_MAP.get(field_name, "global")

        # quantity_kwh: for Nm3, leave None (normalize.py handles conversion)
        quantity_kwh = abs(delta) if unit_raw == "kWh" else None

        records.append(EnergyRecord(
            document_id=f"xls_{uuid.uuid4().hex[:8]}",
            document_type=DocumentType.excel_report,
            source_file=source_file,
            date=date_str,
            supplier="ADWYA_TRIGEN",
            site="STE ADWYA",
            zone=zone,
            energy_type=energy_type,
            quantity_raw=abs(delta),
            unit_raw=unit_raw,
            quantity_kwh=quantity_kwh,
            extraction_confidence=0.95,
            extraction_method="openpyxl_delta",
        ))

    return records


def extract_all_excel(data_dir: str) -> list[EnergyRecord]:
    """Process every .xlsx in data_dir and its reports/ subdirectory."""
    results: list[EnergyRecord] = []
    root = Path(data_dir)
    patterns = list(root.glob("*.xlsx")) + list((root / "reports").glob("*.xlsx"))
    for path in sorted(patterns):
        print(f"  [{path.name}]")
        try:
            records = extract_excel(str(path))
            results.extend(records)
            for r in records:
                print(f"    → {r.energy_type.value} | {r.quantity_raw:.0f} {r.unit_raw} | date={r.date}")
        except Exception as e:
            print(f"    ERROR: {e}")
    return results
